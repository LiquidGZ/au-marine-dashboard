#!/usr/bin/env python3
"""
build_data.py — convert the NEWEST master workbook in data/source/ into the JSON files the site reads.

    python scripts/build_data.py            # auto-picks newest data/source/AMR_Master_*.xlsx
    python scripts/build_data.py path.xlsx  # or an explicit workbook

Outputs (site/data/):
    meta.json, categories.json, projects.json, operators.json, zones.json, kpis.json, series.json,
    breakdowns.json, policy.json, news_sources.json, news_rules.json, geo_sources.json, stats.json, build_info.json

Rules the workbook must follow (see its README sheet): sheet names and header cells are looked up by
name; extra columns are passed through untouched; multi-value cells use ' | '.
"""
import datetime as dt
import glob
import json
import os
import re
import shutil
import sys
from collections import Counter, defaultdict

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  pip install -r scripts/requirements.txt")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC_DIR = os.path.join(ROOT, "data", "source")
OUT_DIR = os.path.join(ROOT, "site", "data")
PREFIX = "AMR_Master_"

CATEGORIES = ["fossil", "renewable", "mineral", "agri", "other"]
STAGES = ["Operational", "Under construction", "Approved / FID", "Planning", "Exploration / appraisal",
          "Suspended", "Decommissioning", "Decommissioned", "Cancelled"]
STATES = ["WA", "NT", "QLD", "NSW", "VIC", "TAS", "SA", "EXT"]
TS_RE = re.compile(r"(\d{8})(?:_(\d{4}))?")


def newest_workbook():
    files = [f for f in glob.glob(os.path.join(SRC_DIR, PREFIX + "*.xlsx")) if not os.path.basename(f).startswith("~$")]
    if not files:
        sys.exit(f"No {PREFIX}*.xlsx found in {SRC_DIR}")

    def key(f):
        m = TS_RE.search(os.path.basename(f))
        return (m.group(1), m.group(2) or "0000") if m else ("0", "0")

    return sorted(files, key=key)[-1]


def rows(ws):
    """Yield dicts keyed by header row; skip fully-empty rows."""
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    for r in it:
        if r is None or all(v in (None, "") for v in r):
            continue
        d = {}
        for h, v in zip(header, r):
            if not h:
                continue
            if isinstance(v, str):
                v = v.strip()
                if v == "":
                    v = None
            if isinstance(v, (dt.datetime, dt.date)):
                v = v.isoformat()[:10]
            d[h] = v
        yield d


def split(v, sep="|"):
    if v is None:
        return []
    if isinstance(v, (int, float)):
        return [str(v)]
    return [s.strip() for s in str(v).split(sep) if s.strip()]


def num(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def intv(v):
    x = num(v)
    return int(x) if x is not None else None


def to_bool(v):
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in ("true", "1", "yes", "y")


def dump(name, obj):
    os.makedirs(OUT_DIR, exist_ok=True)
    p = os.path.join(OUT_DIR, name)
    with open(p, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=1)
    print(f"  wrote {name:22s} {os.path.getsize(p)/1024:7.1f} KB")


def kv(ws):
    out = {}
    for r in rows(ws):
        if r.get("key"):
            out[r["key"]] = r.get("value")
    return out


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else newest_workbook()
    print(f"Building from {os.path.relpath(path, ROOT)}")
    wb = load_workbook(path, read_only=True, data_only=True)
    need = ["Meta", "Categories", "Projects", "Operators", "Zones", "KPIs", "Series", "SeriesData", "Breakdowns",
            "BreakdownItems", "Policy_Stages", "Policy_Jurisdictions", "Policy_Instruments", "Policy_Changes",
            "Policy_Notes", "NewsSources", "NewsRules", "GeoSources"]
    missing = [s for s in need if s not in wb.sheetnames]
    if missing:
        sys.exit(f"Workbook is missing sheets: {missing}")
    warn = []

    # ---- Meta
    meta = kv(wb["Meta"])
    for k in ("map_center_lat", "map_center_lng", "map_zoom", "news_days_to_keep"):
        if k in meta:
            meta[k] = num(meta[k])
    meta.update(kv(wb["Policy_Notes"]))
    dump("meta.json", meta)

    # ---- Categories
    cats = {}
    subs = defaultdict(list)
    for r in rows(wb["Categories"]):
        if not r.get("category_id"):
            continue
        r["order"] = intv(r.get("order")) or 99
        if r.get("subcategory"):
            subs[r["category_id"]].append(r)
        else:
            cats[r["category_id"]] = r
    categories = []
    for cid in CATEGORIES:
        c = cats.get(cid, {"category_id": cid, "name_en": cid, "name_zh": cid})
        c["subcategories"] = sorted(subs.get(cid, []), key=lambda s: s["order"])
        categories.append(c)
    dump("categories.json", categories)
    subnames = {(c["category_id"], s["subcategory"]) for c in categories for s in c["subcategories"]}

    # ---- Projects
    projects = []
    for r in rows(wb["Projects"]):
        if not r.get("project_id"):
            continue
        for k in ("lat", "lng", "capacity", "annual_output", "annual_value_aud_m", "capex_aud_bn"):
            r[k] = num(r.get(k))
        for k in ("start_year", "end_year"):
            r[k] = intv(r.get(k))
        r["sources"] = split(r.get("sources"))
        if r.get("category") not in CATEGORIES:
            warn.append(f"project {r['project_id']}: unknown category '{r.get('category')}'")
        if (r.get("category"), r.get("subcategory")) not in subnames:
            warn.append(f"project {r['project_id']}: sub-category '{r.get('subcategory')}' not in Categories sheet")
        if r.get("stage") not in STAGES:
            warn.append(f"project {r['project_id']}: unknown stage '{r.get('stage')}'")
        if r["lat"] is None or r["lng"] is None:
            warn.append(f"project {r['project_id']}: no coordinates — not on the map")
        projects.append(r)
    dump("projects.json", projects)

    # ---- Operators
    operators = []
    by_op = defaultdict(list)
    for p in projects:
        if p.get("operator"):
            by_op[p["operator"]].append(p)
    for r in rows(wb["Operators"]):
        if not r.get("operator_id"):
            continue
        r["categories"] = split(r.get("categories"))
        r["sources"] = split(r.get("sources"))
        ps = by_op.get(r["name"], [])
        r["project_ids"] = [p["project_id"] for p in ps]
        r["by_stage"] = dict(Counter(p["stage"] for p in ps))
        r["by_category"] = dict(Counter(p["category"] for p in ps))
        operators.append(r)
    known = {o["name"] for o in operators}
    for name in sorted(set(by_op) - known):
        warn.append(f"operator '{name}' used by {len(by_op[name])} project(s) has no Operators row")
    dump("operators.json", operators)

    # ---- Zones
    zones = []
    for r in rows(wb["Zones"]):
        if not r.get("zone_id"):
            continue
        for k in ("area_km2", "potential_gw", "fl_capacity_gw", "centroid_lat", "centroid_lng"):
            r[k] = num(r.get(k))
        for k in ("fl_granted", "fl_active", "commercial_licences"):
            r[k] = intv(r.get(k))
        r["sources"] = split(r.get("sources"))
        zones.append(r)
    dump("zones.json", zones)

    # ---- KPIs
    kpis = []
    for r in rows(wb["KPIs"]):
        if not r.get("kpi_id"):
            continue
        r["value"] = num(r.get("value"))
        r["kpi"] = to_bool(r.get("kpi"))
        r["kpi_order"] = num(r.get("kpi_order"))
        kpis.append(r)
    dump("kpis.json", kpis)

    # ---- Series
    pts = defaultdict(list)
    for r in rows(wb["SeriesData"]):
        if r.get("series_id") and r.get("period") is not None and num(r.get("value")) is not None:
            pts[r["series_id"]].append({"period": str(r["period"]), "value": num(r["value"])})
    series = []
    for r in rows(wb["Series"]):
        if not r.get("series_id"):
            continue
        r["points"] = sorted(pts.get(r["series_id"], []), key=lambda x: x["period"])
        series.append(r)
    dump("series.json", series)

    items = defaultdict(list)
    for r in rows(wb["BreakdownItems"]):
        if r.get("breakdown_id"):
            r["value"] = num(r.get("value"))
            items[r["breakdown_id"]].append(r)
    breakdowns = []
    for r in rows(wb["Breakdowns"]):
        if r.get("breakdown_id"):
            r["items"] = items.get(r["breakdown_id"], [])
            breakdowns.append(r)
    dump("breakdowns.json", breakdowns)

    # ---- Policy
    stages = sorted([r for r in rows(wb["Policy_Stages"]) if r.get("stage_id")], key=lambda s: intv(s.get("order")) or 99)
    stage_ids = {s["stage_id"] for s in stages}
    instruments = []
    for r in rows(wb["Policy_Instruments"]):
        if not r.get("policy_id"):
            continue
        r["categories"] = split(r.get("categories"))
        r["stages"] = split(r.get("stages"))
        r["year"] = intv(r.get("year"))
        bad = [s for s in r["stages"] if s not in stage_ids]
        if bad:
            warn.append(f"policy {r['policy_id']}: unknown stage(s) {bad}")
        instruments.append(r)
    changes = [r for r in rows(wb["Policy_Changes"]) if r.get("jurisdiction_id")]
    for c in changes:
        c["categories"] = split(c.get("categories"))
        c["date"] = str(c.get("date") or "")
    changes.sort(key=lambda c: c["date"], reverse=True)
    jurisdictions = [r for r in rows(wb["Policy_Jurisdictions"]) if r.get("jurisdiction_id")]
    dump("policy.json", {"stages": stages, "jurisdictions": jurisdictions, "instruments": instruments, "changes": changes})

    # ---- News + geo config
    news_sources = []
    for r in rows(wb["NewsSources"]):
        if r.get("url"):
            r["enabled"] = to_bool(r.get("enabled", True))
            r["filter_keywords"] = split(r.get("filter_keywords"), ",")
            news_sources.append(r)
    dump("news_sources.json", news_sources)
    rules = []
    for r in rows(wb["NewsRules"]):
        if r.get("category"):
            r["keywords_en"] = [k.lower() for k in split(r.get("keywords_en"), ",")]
            r["keywords_zh"] = split(r.get("keywords_zh"), ",")
            rules.append(r)
    dump("news_rules.json", rules)
    geo = []
    for r in rows(wb["GeoSources"]):
        if r.get("layer_id"):
            r["default_on"] = to_bool(r.get("default_on"))
            r["enabled"] = to_bool(r.get("enabled", True))
            geo.append(r)
    dump("geo_sources.json", geo)

    # ---- Derived statistics
    cat_stage = {c: {s: 0 for s in STAGES} for c in CATEGORIES}
    sub_stage = defaultdict(lambda: {s: 0 for s in STAGES})
    state_cat = {s: {c: 0 for c in CATEGORIES} for s in STATES}
    for p in projects:
        if p.get("category") in cat_stage and p.get("stage") in STAGES:
            cat_stage[p["category"]][p["stage"]] += 1
            sub_stage[f'{p["category"]}|{p.get("subcategory")}'][p["stage"]] += 1
        if p.get("state") in state_cat and p.get("category") in CATEGORIES:
            state_cat[p["state"]][p["category"]] += 1

    # capacity by stage: for each sub-category use its dominant unit, so totals are like-for-like
    cap = []
    groups = defaultdict(list)
    for p in projects:
        if p.get("capacity") is not None and p.get("capacity_unit"):
            groups[(p["category"], p.get("subcategory"))].append(p)
    for (cat, sub), ps in groups.items():
        unit = Counter(p["capacity_unit"] for p in ps).most_common(1)[0][0]
        same = [p for p in ps if p["capacity_unit"] == unit]
        if len(same) < 3:
            continue
        by = {s: 0.0 for s in STAGES}
        for p in same:
            by[p["stage"]] = round(by[p["stage"]] + p["capacity"], 3)
        cap.append({"category": cat, "subcategory": sub, "unit": unit, "n": len(same), "excluded": len(ps) - len(same), "by_stage": by})
    cap.sort(key=lambda c: (CATEGORIES.index(c["category"]), c["subcategory"] or ""))

    capex = {c: {s: 0.0 for s in STAGES} for c in CATEGORIES}
    capex_n = Counter()
    for p in projects:
        if p.get("capex_aud_bn"):
            capex[p["category"]][p["stage"]] = round(capex[p["category"]][p["stage"]] + p["capex_aud_bn"], 3)
            capex_n[p["category"]] += 1

    def ym(s):
        m = re.match(r"(\d{4})-(\d{2})", str(s or ""))
        return f"{m.group(1)}-{m.group(2)}" if m else None
    updates = sorted([{"project_id": p["project_id"], "date": ym(p.get("latest_update_en"))} for p in projects if ym(p.get("latest_update_en"))],
                     key=lambda u: u["date"], reverse=True)

    dump("stats.json", {
        "project_count": len(projects),
        "by_category_stage": cat_stage,
        "by_subcategory_stage": sub_stage,
        "by_state_category": state_cat,
        "by_stage": {s: sum(cat_stage[c][s] for c in CATEGORIES) for s in STAGES},
        "capacity_by_stage": cap,
        "capex_by_category_stage": capex, "capex_counts": capex_n,
        "latest_updates": updates[:40],
        "operator_count": len(operators),
        "instrument_count": len(instruments),
    })

    # copy the workbook next to the JSON so the site can offer it for download
    wb_dir = os.path.join(OUT_DIR, "workbook")
    os.makedirs(wb_dir, exist_ok=True)
    for old in glob.glob(os.path.join(wb_dir, "*.xlsx")):
        os.remove(old)
    shutil.copy2(path, os.path.join(wb_dir, os.path.basename(path)))

    dump("build_info.json", {
        "workbook": os.path.basename(path),
        "built_at_utc": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "counts": {"projects": len(projects), "operators": len(operators), "zones": len(zones), "kpis": len(kpis),
                   "series": len(series), "policy_instruments": len(instruments), "policy_changes": len(changes),
                   "news_sources": len(news_sources)},
        "warnings": warn,
    })
    for w in warn:
        print("  WARN " + w)
    print(f"Done: {len(projects)} projects, {len(operators)} operators, {len(instruments)} policy instruments, {len(warn)} warnings")


if __name__ == "__main__":
    main()
