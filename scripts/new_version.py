#!/usr/bin/env python3
"""
new_version.py — start a new time-stamped version of the master workbook.

    python scripts/new_version.py            # timestamp = today (YYYYMMDD); adds _HHMM if a file for today exists
    python scripts/new_version.py "notes"    # optional note appended to the Changelog sheet of the NEW copy

What it does
  1. finds the newest data/source/AMR_Master_*.xlsx
  2. copies it to data/archive/ (unchanged, as the frozen older version)
  3. creates data/source/AMR_Master_<timestamp>.xlsx — the copy you now edit
  4. removes the older file from data/source/ so build_data.py always picks the newest one unambiguously
  5. appends a Changelog row to the new copy
Then: open the new file in Excel, edit, save, commit and push.
"""
import datetime as dt
import glob
import os
import re
import shutil
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "data", "source")
ARC = os.path.join(ROOT, "data", "archive")
TS_RE = re.compile(r"(\d{8})(?:_(\d{4}))?")


def newest():
    files = [f for f in glob.glob(os.path.join(SRC, "AMR_Master_*.xlsx")) if not os.path.basename(f).startswith("~$")]
    if not files:
        sys.exit("No AMR_Master_*.xlsx in data/source/")
    return sorted(files, key=lambda f: (TS_RE.search(f).group(1), TS_RE.search(f).group(2) or "0000"))[-1]


def main():
    note = " ".join(sys.argv[1:]).strip()
    cur = newest()
    today = dt.date.today().strftime("%Y%m%d")
    new = os.path.join(SRC, f"AMR_Master_{today}.xlsx")
    if os.path.abspath(new) == os.path.abspath(cur) or os.path.exists(new):
        new = os.path.join(SRC, f"AMR_Master_{today}_{dt.datetime.now().strftime('%H%M')}.xlsx")
    os.makedirs(ARC, exist_ok=True)
    archived = os.path.join(ARC, os.path.basename(cur))
    shutil.copy2(cur, archived)
    shutil.copy2(cur, new)
    os.remove(cur)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(new)
        if "Changelog" in wb.sheetnames:
            wb["Changelog"].append([os.path.basename(new), dt.date.today().isoformat(), os.environ.get("USERNAME") or os.environ.get("USER") or "",
                                    note or "New version created with scripts/new_version.py"])
            wb.save(new)
    except Exception as e:  # noqa: BLE001
        print(f"(could not append Changelog row: {e})")
    print(f"archived : {os.path.relpath(archived, ROOT)}")
    print(f"new      : {os.path.relpath(new, ROOT)}   <- edit this one, then commit + push")


if __name__ == "__main__":
    main()
